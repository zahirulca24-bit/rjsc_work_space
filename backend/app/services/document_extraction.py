import csv
import io
from pathlib import Path
from typing import Optional

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader
import xlrd


TEXT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".xlsx",
    ".xls",
    ".csv",
}

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
}


class UnsupportedExtractionError(Exception):
    pass


def _clean_text(value: str, max_chars: int) -> str:
    value = value.replace("\x00", " ")
    value = "\n".join(
        line.strip()
        for line in value.splitlines()
        if line.strip()
    )
    return value[:max_chars]


def extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))

    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text:
            parts.append(text)

    return "\n".join(parts)


def extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))

    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]
            parts.append(" | ".join(values))

    return "\n".join(parts)


def extract_xlsx(data: bytes) -> str:
    workbook = load_workbook(
        io.BytesIO(data),
        read_only=True,
        data_only=True,
    )

    parts = []

    for sheet in workbook.worksheets:
        parts.append(f"[Sheet: {sheet.title}]")

        for row in sheet.iter_rows(values_only=True):
            values = [
                "" if value is None else str(value)
                for value in row
            ]

            if any(value.strip() for value in values):
                parts.append(" | ".join(values))

    workbook.close()
    return "\n".join(parts)


def extract_xls(data: bytes) -> str:
    workbook = xlrd.open_workbook(
        file_contents=data,
    )

    parts = []

    for sheet in workbook.sheets():
        parts.append(f"[Sheet: {sheet.name}]")

        for row_index in range(sheet.nrows):
            values = [
                str(sheet.cell_value(row_index, col))
                for col in range(sheet.ncols)
            ]

            if any(value.strip() for value in values):
                parts.append(" | ".join(values))

    return "\n".join(parts)


def extract_csv(data: bytes) -> str:
    decoded = data.decode(
        "utf-8-sig",
        errors="replace",
    )

    rows = csv.reader(io.StringIO(decoded))

    return "\n".join(
        " | ".join(cell.strip() for cell in row)
        for row in rows
    )


def extract_document_text(
    data: bytes,
    filename: Optional[str],
    max_chars: int,
) -> str:
    extension = Path(filename or "").suffix.lower()

    if extension == ".pdf":
        text = extract_pdf(data)

    elif extension == ".docx":
        text = extract_docx(data)

    elif extension == ".xlsx":
        text = extract_xlsx(data)

    elif extension == ".xls":
        text = extract_xls(data)

    elif extension == ".csv":
        text = extract_csv(data)

    elif extension in IMAGE_EXTENSIONS:
        raise UnsupportedExtractionError(
            "Image documents require vision analysis"
        )

    else:
        raise UnsupportedExtractionError(
            f"Unsupported extraction type: {extension}"
        )

    return _clean_text(text, max_chars)
